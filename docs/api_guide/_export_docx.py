# -*- coding: utf-8 -*-
"""
Build a well-formatted Customer SDK API Guide as .docx from hand-written markdown.

Design choices (Word readability):
- Reading order matches STRUCTURE.md (secondary-dev → AI → guides → tech chapters)
- Strip MD-only nav footers / repeated canonical notes (keep once in front matter)
- Heading 1–3 map to Word styles; page break before each H1 chapter
- Tables → real Word tables with header shading
- Fenced code → Consolas + light gray cell (not raw ```)
- Inline `code`, **bold**, *italic* preserved
- CJK body: 微软雅黑; Western: Calibri; code: Consolas
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "export"

# Chapter order: (relative path under lang dir, optional section label for TOC grouping)
CHAPTERS = [
    ("00_secondary_dev.md", "入门 / Getting started"),
    ("AI_DEV.md", "入门 / Getting started"),
    ("AI_CONTEXT.md", "入门 / Getting started"),
    ("guides/G00_directions.md", "场景指南 / Guides"),
    ("guides/G01_modbus_rs485_lvgl.md", "场景指南 / Guides"),
    ("guides/G02_wifi_mqtt_panel.md", "场景指南 / Guides"),
    ("guides/G03_ble_local.md", "场景指南 / Guides"),
    ("guides/G04_eth_rs485_gateway.md", "场景指南 / Guides"),
    ("guides/G05_full_fusion.md", "场景指南 / Guides"),
    ("guides/G06_ai_prompts.md", "场景指南 / Guides"),
    ("01_overview.md", "技术分章 / Technical"),
    ("02_build.md", "技术分章 / Technical"),
    ("03_boot.md", "技术分章 / Technical"),
    ("04_wifi.md", "技术分章 / Technical"),
    ("05_mqtt.md", "技术分章 / Technical"),
    ("06_bt.md", "技术分章 / Technical"),
    ("07_eth.md", "技术分章 / Technical"),
    ("08_backlight.md", "技术分章 / Technical"),
    ("09_beep.md", "技术分章 / Technical"),
    ("10_gpio_rs485.md", "技术分章 / Technical"),
    ("11_modbus.md", "技术分章 / Technical"),
    ("12_coexist.md", "技术分章 / Technical"),
    ("13_low_power.md", "技术分章 / Technical"),
    ("14_gui_task.md", "技术分章 / Technical"),
    ("15_hub.md", "技术分章 / Technical"),
    ("16_icons.md", "技术分章 / Technical"),
    ("17_health.md", "技术分章 / Technical"),
    ("18_kconfig.md", "技术分章 / Technical"),
    ("19_control_map.md", "技术分章 / Technical"),
    ("20_faq.md", "技术分章 / Technical"),
    ("A_headers.md", "技术分章 / Technical"),
]

NAV_LINE_RE = re.compile(
    r"^(\[←[^\]]*\]\([^)]*\)\s*\|\s*)?\[(目录|TOC|README)\]\([^)]*\)"
    r"|^\[[^\]]+→\]\([^)]*\)"
    r"|^\[←[^\]]*\]\([^)]*\)\s*\|\s*\[(目录|TOC|README)\]"
)
FOOTER_NOTE_RE = re.compile(
    r"^>\s*(权威定义|Authoritative definitions|Canonical:)", re.I
)
HR_RE = re.compile(r"^---+\s*$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")
FENCE_RE = re.compile(r"^```(\w*)\s*$")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")
UL_RE = re.compile(r"^(\s*)[-*+]\s+(.+)$")
OL_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.+)$")
BQ_RE = re.compile(r"^>\s?(.*)$")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
INLINE_CODE_RE = re.compile(r"`([^`]+)`")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
ITALIC_RE = re.compile(r"(?<!\*)\*([^*]+?)\*(?!\*)")


def set_run_font(run, east_asia: str, ascii_font: str, size_pt: float | None = None, bold=None):
    run.font.name = ascii_font
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east_asia)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def set_paragraph_spacing(p, before=6, after=6, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    """kwargs: top/left/bottom/right -> {sz, color, val}"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            el = OxmlElement(f"w:{edge}")
            for k, v in kwargs[edge].items():
                el.set(qn(f"w:{k}"), str(v))
            tcBorders.append(el)
    tcPr.append(tcBorders)


def add_page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, "微软雅黑", "Calibri", 9)

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run2 = p.add_run()
    run2._r.append(fld_char_begin)
    run2._r.append(instr)
    run2._r.append(fld_char_end)
    set_run_font(run2, "微软雅黑", "Calibri", 9)


def configure_styles(doc: Document, lang: str):
    styles = doc.styles
    body_ea = "微软雅黑" if lang == "zh" else "Microsoft YaHei"
    body_ascii = "Calibri"

    normal = styles["Normal"]
    normal.font.name = body_ascii
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), body_ea)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(6)

    for level, size, before, after in (
        (1, 18, 18, 10),
        (2, 14, 14, 8),
        (3, 12, 10, 6),
    ):
        style = styles[f"Heading {level}"]
        style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.name = body_ascii
        style._element.rPr.rFonts.set(qn("w:eastAsia"), body_ea)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_inline_runs(paragraph, text: str, lang: str, base_size=11, code=False):
    """Parse inline markdown into runs."""
    ea = "微软雅黑" if lang == "zh" else "Microsoft YaHei"
    # Rewrite links to visible text only (drop ./paths)
    text = LINK_RE.sub(r"\1", text)

    # Tokenize: **bold**, `code`, *italic*
    pattern = re.compile(
        r"(\*\*[^*]+?\*\*|`[^`]+`|\*[^*]+?\*)"
    )
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos : m.start()])
            set_run_font(run, ea, "Consolas" if code else "Calibri", base_size)
        token = m.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, ea, "Calibri", base_size, bold=True)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, ea, "Consolas", base_size - 0.5)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, ea, "Calibri", base_size)
            run.italic = True
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, ea, "Consolas" if code else "Calibri", base_size)


def should_skip_line(line: str) -> bool:
    s = line.strip()
    if not s:
        return False
    if NAV_LINE_RE.search(s):
        return True
    if FOOTER_NOTE_RE.search(s):
        return True
    # Pure nav combo lines
    if s.startswith("[←") and "](./" in s and "|" in s:
        return True
    if s in ("[← README](./README.md) | [Environment & build →](./02_build.md)",):
        return True
    return False


def clean_md(text: str) -> str:
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    out = []
    for line in lines:
        if should_skip_line(line):
            continue
        out.append(line)
    # Collapse 3+ blank lines
    cleaned = []
    blank = 0
    for line in out:
        if not line.strip():
            blank += 1
            if blank <= 2:
                cleaned.append("")
        else:
            blank = 0
            cleaned.append(line)
    return "\n".join(cleaned).strip() + "\n"


def parse_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def add_code_block(doc: Document, code: str, lang: str):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, "F5F5F5")
    border = {"sz": "4", "val": "single", "color": "DDDDDD"}
    set_cell_border(cell, top=border, left=border, bottom=border, right=border)

    # Clear default paragraph text
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    set_paragraph_spacing(p, before=4, after=2, line=1.05)
    lines = code.rstrip("\n").split("\n") or [""]
    for i, line in enumerate(lines):
        if i == 0:
            target = p
        else:
            target = cell.add_paragraph()
            set_paragraph_spacing(target, before=0, after=0, line=1.05)
        run = target.add_run(line if line else " ")
        set_run_font(run, "微软雅黑", "Consolas", 9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    # Spacer after code
    sp = doc.add_paragraph()
    set_paragraph_spacing(sp, before=2, after=2)


def add_table(doc: Document, rows: list[list[str]], lang: str):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    rows = [r + [""] * (cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, before=3, after=3, line=1.1)
            # Strip markdown links inside cells
            plain = LINK_RE.sub(r"\1", text)
            add_inline_runs(p, plain, lang, base_size=9.5)
            if i == 0:
                shade_cell(cell, "E8EEF7")
                for run in p.runs:
                    run.bold = True
    sp = doc.add_paragraph()
    set_paragraph_spacing(sp, before=4, after=4)


def add_heading(doc: Document, text: str, level: int, lang: str, first_h1: list):
    # Word Heading styles are 1..9; clamp
    level = min(max(level, 1), 3)
    text = LINK_RE.sub(r"\1", text.strip())
    # Drop leading "# " already stripped
    if level == 1 and not first_h1[0]:
        first_h1[0] = True
    elif level == 1:
        doc.add_page_break()
    p = doc.add_heading(text, level=level)
    ea = "微软雅黑" if lang == "zh" else "Microsoft YaHei"
    for run in p.runs:
        set_run_font(run, ea, "Calibri", None, bold=True)


def render_md(doc: Document, md: str, lang: str, first_h1: list):
    lines = md.split("\n")
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if HR_RE.match(stripped):
            # Visual separator as thin paragraph
            p = doc.add_paragraph("─" * 24)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_run_font(run, "微软雅黑", "Calibri", 9)
                run.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
            set_paragraph_spacing(p, before=8, after=8)
            i += 1
            continue

        m = FENCE_RE.match(stripped)
        if m:
            i += 1
            buf = []
            while i < n and not FENCE_RE.match(lines[i].strip()):
                buf.append(lines[i])
                i += 1
            if i < n:
                i += 1  # closing fence
            add_code_block(doc, "\n".join(buf), lang)
            continue

        hm = HEADING_RE.match(stripped)
        if hm:
            add_heading(doc, hm.group(2), len(hm.group(1)), lang, first_h1)
            i += 1
            continue

        # Table block
        if stripped.startswith("|") and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1].strip()):
            rows = [parse_table_row(stripped)]
            i += 2  # skip header + separator
            while i < n and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i].strip()))
                i += 1
            add_table(doc, rows, lang)
            continue

        bq = BQ_RE.match(line)
        if bq:
            buf = [bq.group(1)]
            i += 1
            while i < n:
                b2 = BQ_RE.match(lines[i])
                if not b2:
                    break
                buf.append(b2.group(1))
                i += 1
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=6, after=6, line=1.2)
            p.paragraph_format.left_indent = Cm(0.5)
            # Left bar simulation via prefix
            run = p.add_run("▎ ")
            set_run_font(run, "微软雅黑", "Calibri", 11)
            run.font.color.rgb = RGBColor(0x2F, 0x6F, 0xAD)
            add_inline_runs(p, " ".join(buf), lang, base_size=10.5)
            for run in p.runs[1:]:
                run.italic = True
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            continue

        um = UL_RE.match(line)
        if um:
            text = um.group(2)
            p = doc.add_paragraph(style="List Bullet")
            set_paragraph_spacing(p, before=2, after=2)
            add_inline_runs(p, text, lang)
            i += 1
            continue

        om = OL_RE.match(line)
        if om:
            text = om.group(3)
            p = doc.add_paragraph(style="List Number")
            set_paragraph_spacing(p, before=2, after=2)
            add_inline_runs(p, text, lang)
            i += 1
            continue

        # Normal paragraph (merge continuation lines that aren't special)
        buf = [stripped]
        i += 1
        while i < n:
            nxt = lines[i]
            ns = nxt.strip()
            if (
                not ns
                or HR_RE.match(ns)
                or FENCE_RE.match(ns)
                or HEADING_RE.match(ns)
                or ns.startswith("|")
                or BQ_RE.match(nxt)
                or UL_RE.match(nxt)
                or OL_RE.match(nxt)
            ):
                break
            buf.append(ns)
            i += 1
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=4, after=6, line=1.25)
        add_inline_runs(p, " ".join(buf), lang)


def add_cover(doc: Document, lang: str):
    for _ in range(2):
        doc.add_paragraph()

    title = (
        "Customer SDK\n二次开发 API 手册"
        if lang == "zh"
        else "Customer SDK\nSecondary-Development API Guide"
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title.split("\n")[0])
    set_run_font(run, "微软雅黑", "Calibri", 28, bold=True)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(title.split("\n")[1])
    set_run_font(run2, "微软雅黑", "Calibri", 22, bold=True)
    run2.font.color.rgb = RGBColor(0x2F, 0x6F, 0xAD)

    doc.add_paragraph()
    sub = (
        "开放层 API · 场景指南 · AI 协作约束\n面向产品二次开发（不含封闭库反汇编）"
        if lang == "zh"
        else "Open-layer APIs · Scenario guides · AI constraints\nProduct secondary development only"
    )
    for line in sub.split("\n"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        set_run_font(run, "微软雅黑", "Calibri", 12)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()
    meta = (
        "建议阅读顺序：二次开发入门 →（可选）AI 协作 → G00 方向 → G01 实战 → 技术分章按需"
        if lang == "zh"
        else "Suggested order: Getting started → (optional) AI-assisted → G00 → G01 → technical chapters as needed"
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Cm(1.5)
    p.paragraph_format.right_indent = Cm(1.5)
    run = p.add_run(meta)
    set_run_font(run, "微软雅黑", "Calibri", 10.5)
    run.italic = True

    note = (
        "权威定义以工程内头文件为准（include/*.h、main/app/*.h）。封闭库（.a）请勿反汇编。"
        if lang == "zh"
        else "Authoritative definitions: include/*.h and main/app/*.h. Do not reverse-engineer closed .a libraries."
    )
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(note)
    set_run_font(run, "微软雅黑", "Calibri", 9)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_page_break()


def add_toc_placeholder(doc: Document, lang: str):
    title = "目录" if lang == "zh" else "Table of Contents"
    p = doc.add_heading(title, level=1)
    ea = "微软雅黑"
    for run in p.runs:
        set_run_font(run, ea, "Calibri", None, bold=True)

    hint = (
        "（在 Word 中：引用 → 目录 → 自动目录；或选中下方域后按 F9 更新）"
        if lang == "zh"
        else "(In Word: References → Table of Contents; or select the field below and press F9)"
    )
    hp = doc.add_paragraph()
    run = hp.add_run(hint)
    set_run_font(run, ea, "Calibri", 9)
    run.italic = True
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    # TOC field
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")

    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = r' TOC \o "1-2" \h \z \u '

    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")

    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")

    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run2 = paragraph.add_run(
        "右键此处 → 更新域" if lang == "zh" else "Right-click → Update field"
    )
    set_run_font(run2, ea, "Calibri", 10)
    run2._r.append(fldChar3)

    doc.add_page_break()


def add_part_banner(doc: Document, label: str, lang: str, seen: set, first_h1: list) -> None:
    """Section divider: always on a fresh page so it is not orphaned before a chapter break."""
    part = label.split("/")[0].strip() if lang == "zh" else label.split("/")[-1].strip()
    if part in seen:
        return
    if seen:
        doc.add_page_break()
    seen.add(part)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, before=4, after=10)
    run = p.add_run("■  " + part)
    set_run_font(run, "微软雅黑", "Calibri", 14, bold=True)
    run.font.color.rgb = RGBColor(0x2F, 0x6F, 0xAD)
    # Next chapter H1 should not insert another page break
    first_h1[0] = False


def build(lang: str) -> Path:
    base = ROOT / lang
    if not base.is_dir():
        raise SystemExit(f"Missing language dir: {base}")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    add_page_number_footer(section)
    configure_styles(doc, lang)

    add_cover(doc, lang)
    add_toc_placeholder(doc, lang)

    # After TOC page break: first chapter H1 must not insert another break.
    first_h1 = [False]
    seen_parts: set[str] = set()
    missing = []
    for rel, part_label in CHAPTERS:
        path = base / rel
        if not path.is_file():
            missing.append(str(path))
            continue
        add_part_banner(doc, part_label, lang, seen_parts, first_h1)
        md = clean_md(path.read_text(encoding="utf-8"))
        render_md(doc, md, lang, first_h1)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    name = (
        "Customer_SDK_API_Guide_ZH.docx"
        if lang == "zh"
        else "Customer_SDK_API_Guide_EN.docx"
    )
    out = OUT_DIR / name
    doc.save(out)
    if missing:
        print("WARN missing:", *missing, sep="\n  ")
    return out


def main():
    langs = sys.argv[1:] or ["zh", "en"]
    for lang in langs:
        out = build(lang)
        print(f"Wrote {out} ({out.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
